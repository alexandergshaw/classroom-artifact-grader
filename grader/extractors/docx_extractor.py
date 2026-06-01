from __future__ import annotations

from pathlib import Path

from docx import Document


def extract_docx(path: str | Path) -> tuple[str, list[str]]:
    doc = Document(str(path))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    headings = [
        p.text.strip()
        for p in doc.paragraphs
        if p.text.strip() and p.style and str(p.style.name).lower().startswith("heading")
    ]
    return "\n".join(paragraphs), headings
